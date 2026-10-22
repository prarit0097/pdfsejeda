"""Live end-to-end check against a running server: every tool, through the real HTTP API.

    python scripts/smoke_test.py [base_url]        (default http://127.0.0.1:8000)

For each tool it uploads sample files, POSTs /api/tasks, polls /change, downloads the
result and checks it is the expected kind of file. Exit code 1 if anything fails.
"""
from __future__ import annotations

import io
import sys
import time
import zipfile

import httpx
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas

BASE = (sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8000").rstrip("/")
c = httpx.Client(base_url=BASE, timeout=180)


# ---------------------------------------------------------------- sample files

def pdf(pages=3, size=A4, bookmarks=False, texts=None, table=False, link=False, image=False):
    buf = io.BytesIO()
    cv = canvas.Canvas(buf, pagesize=size)
    cv.setTitle("Smoke test")
    for i in range(pages):
        w, h = size
        cv.setFont("Helvetica-Bold", 22)
        cv.drawString(72, h - 72, texts[i] if texts else f"Page {i + 1}")
        cv.setFont("Helvetica", 12)
        for k in range(12):
            cv.drawString(72, h - 120 - k * 18, "The quick brown fox jumps over the lazy dog. 0123456789")
        if table and i == 0:
            rows = [["Item", "Qty", "Price"], ["Apples", "3", "1.50"], ["Pears", "7", "2.25"]]
            for r, row in enumerate(rows):
                for col, v in enumerate(row):
                    cv.rect(72 + col * 110, 300 - r * 22, 110, 22)
                    cv.drawString(78 + col * 110, 307 - r * 22, v)
        if link and i == 0:
            cv.linkURL("https://example.com", (72, 60, 200, 80))
            cv.drawString(72, 66, "example.com")
        if image and i == 0:
            from PIL import Image
            from reportlab.lib.utils import ImageReader
            import numpy as np

            noise = (np.random.default_rng(i).random((900, 900, 3)) * 255).astype("uint8")
            cv.drawImage(ImageReader(Image.fromarray(noise)), 300, 300, 250, 250)
        if bookmarks:
            cv.bookmarkPage(f"p{i}")
            cv.addOutlineEntry(f"Chapter {i + 1}", f"p{i}", level=0)
        cv.showPage()
    cv.save()
    return buf.getvalue()


def png(w=300, h=200):
    from PIL import Image, ImageDraw

    im = Image.new("RGB", (w, h), "white")
    ImageDraw.Draw(im).rectangle([20, 20, w - 20, h - 20], outline="navy", width=6)
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def scan_pdf():
    """A page that is only an image of text, for OCR."""
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(pdf(1, texts=["Scanned Invoice 4711"]))
    pil = doc[0].render(scale=300 / 72).to_pil().convert("L")
    b = io.BytesIO()
    pil.save(b, "PDF", resolution=300)
    return b.getvalue()


def docx():
    from docx import Document

    d = Document()
    d.add_heading("Smoke test letter", 1)
    d.add_paragraph("This paragraph came from Word.")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def form_pdf():
    """A PDF with a text field and a checkbox, made by our own form creator."""
    fid = upload("blank-form.pdf", pdf(1))
    r = run("editForms", [fid], appendFormFieldOperations=[
        {"page": "1", "kind": "text", "name": "name", "boundingBox": {"left": 72, "bottom": 600, "right": 300, "top": 620}},
        {"page": "1", "kind": "checkbox", "name": "agree", "boundingBox": {"left": 72, "bottom": 560, "right": 86, "top": 574}}])
    return r["data"]


# ---------------------------------------------------------------- HTTP helpers

def upload(name, data, mime="application/octet-stream"):
    r = c.post("/api/files/upload", files={"file": (name, data, mime)})
    r.raise_for_status()
    return r.json()["id"]


def run(task_type, files, **params):
    r = c.post("/api/tasks", json={"type": task_type, "files": files, **params})
    if r.status_code != 200:
        raise AssertionError(f"HTTP {r.status_code}: {r.text[:200]}")
    tid = r.json()
    end = time.time() + 240
    while time.time() < end:
        t = c.get(f"/api/tasks/{tid}/change").json()
        if t["status"] in ("Completed", "Failed"):
            break
        time.sleep(0.4)
    if t["status"] != "Completed":
        raise AssertionError(t.get("failureReason") or t["status"])
    url = "/" + t["resultFullUrlPattern"].split("/", 1)[1].replace("%7B%7Bfilename%7D%7D", t["resultFilename"])
    d = c.get(url)
    d.raise_for_status()
    return {"task": t, "data": d.content, "name": t["resultFilename"]}


def is_pdf(b):
    return b[:5] == b"%PDF-"


def pdf_text(b):
    import pypdf

    return "\n".join(p.extract_text() or "" for p in pypdf.PdfReader(io.BytesIO(b)).pages)


def pdf_pages(b):
    import pypdf

    return len(pypdf.PdfReader(io.BytesIO(b)).pages)


def zip_names(b):
    return zipfile.ZipFile(io.BytesIO(b)).namelist()


# ---------------------------------------------------------------- the checks

def main():
    results = []

    def check(label, fn):
        t0 = time.time()
        try:
            note = fn() or ""
            results.append((True, label, note, time.time() - t0))
            print(f"  PASS  {label:<34} {note}")
        except Exception as e:  # noqa: BLE001
            results.append((False, label, str(e)[:160], time.time() - t0))
            print(f"  FAIL  {label:<34} {str(e)[:160]}")

    print(f"Smoke test against {BASE}\n")
    for path in ["/", "/merge-pdf", "/pdf-editor", "/sign-pdf", "/pdf-forms", "/workflows", "/help", "/privacy"]:
        check(f"page {path}", lambda p=path: (c.get(p).raise_for_status(), "200")[1])
    check("page 404", lambda: "404" if c.get("/nope-nope").status_code == 404 else (_ for _ in ()).throw(AssertionError("not 404")))

    p3 = upload("three.pdf", pdf(3))
    p2 = upload("two.pdf", pdf(2, texts=["B1", "B2"]))
    book = upload("book.pdf", pdf(3, bookmarks=True))
    wide = upload("wide.pdf", pdf(2, landscape(A4)))
    inv = upload("invoices.pdf", pdf(4, texts=["Invoice 100", "Invoice 100", "Invoice 200", "Invoice 300"]))
    tab = upload("table.pdf", pdf(1, table=True))
    img_pdf = upload("photo.pdf", pdf(1, image=True))
    linked = upload("linked.pdf", pdf(1, link=True))
    pic = upload("pic.png", png(), "image/png")

    def expect(cond, msg):
        if not cond:
            raise AssertionError(msg)

    # organize
    check("merge", lambda: (lambda r: (expect(pdf_pages(r["data"]) == 6, "pages"), f"{pdf_pages(r['data'])} pages")[1])(
        run("merge", [p3, p2, pic], tocPolicy="none")))
    check("merge + TOC + footer", lambda: (lambda r: f"{pdf_pages(r['data'])} pages")(
        run("merge", [p3, p2], tocPolicy="file_names", filenameFooter=True, blankPageIfOdd=True)))
    check("split every page", lambda: (lambda r: (expect(len(zip_names(r["data"])) == 3, "files"), "3 files")[1])(
        run("splitByPages", [p3], splitMode="every", step=1)))
    check("split after pages", lambda: (lambda r: f"{len(zip_names(r['data']))} files")(run("splitByPages", [p3], splitMode="after", pages="1")))
    check("split by bookmarks", lambda: (lambda r: ", ".join(zip_names(r["data"])))(run("splitByBookmarks", [book], bookmarksLevel=1)))
    check("split by size", lambda: (lambda r: r["name"])(run("splitBySize", [p3], size=20, sizeUnit="KB")))
    check("split by text", lambda: (lambda r: ", ".join(zip_names(r["data"])))(
        run("splitByText", [inv], area={"x": 60, "y": 40, "width": 300, "height": 50})))
    check("split down the middle", lambda: (lambda r: f"{pdf_pages(r['data'])} pages")(run("splitDownTheMiddle", [wide], mode="auto")))
    check("alternate mix", lambda: (lambda r: f"{pdf_pages(r['data'])} pages")(run("alternateMix", [p3, p2], reverse=True)))
    check("extract pages", lambda: (lambda r: (expect(pdf_pages(r["data"]) == 2, "pages"), "2 pages")[1])(run("extractPages", [p3], pages="1,3")))
    check("delete pages", lambda: (lambda r: (expect(pdf_pages(r["data"]) == 2, "pages"), "2 pages")[1])(run("deletePages", [p3], pages="2")))
    check("organize", lambda: (lambda r: f"{pdf_pages(r['data'])} pages")(run("combineReorder", [p3, p2], filePages=["1:2:0", "0:1:90", "-1:-1:0"])))
    check("rotate", lambda: (lambda r: "ok")(run("rotatePages", [p3], rotation=90)))
    check("rename by text", lambda: (lambda r: r["name"])(run("renameByText", [inv], area={"x": 60, "y": 40, "width": 300, "height": 50})))
    # optimize & layout
    check("compress", lambda: (lambda r: r["task"].get("remark", ""))(run("compress", [img_pdf], imageQuality=0.65, imageDpi=72)))
    check("crop auto", lambda: (lambda r: "ok")(run("crop", [p3], cropMode="auto")))
    check("crop margins", lambda: (lambda r: "ok")(run("crop", [p3], cropMode="margins", unit="cm", marginTop=1, marginBottom=1, marginLeft=1, marginRight=1)))
    check("resize", lambda: (lambda r: "ok")(run("resizePages", [p3], pageSize="letter")))
    check("n-up", lambda: (lambda r: f"{pdf_pages(r['data'])} sheets")(run("nup", [p3], n=4)))
    check("flip", lambda: (lambda r: "ok")(run("flip", [p3], flipMode="horizontal")))
    check("grayscale", lambda: (lambda r: "ok")(run("grayscale", [img_pdf])))
    check("deskew", lambda: (lambda r: r["task"].get("remark", ""))(run("deskew", [p3])))
    check("flatten forms", lambda: (lambda r: "ok")(run("flatten", [p3], flattenMode="onlyforms")))
    check("flatten to images", lambda: (lambda r: "ok")(run("flatten", [p3], flattenMode="all", resolution=75)))
    check("repair", lambda: (lambda r: r["task"].get("remark", ""))(run("repair", [upload("broken.pdf", pdf(2)[:-200])])))
    # stamp
    check("page numbers", lambda: (lambda r: (expect("2 / 3" in pdf_text(r["data"]), "text"), "found '2 / 3'")[1])(
        run("pageNumbers", [p3], pattern="[PAGE_ARABIC] / [TOTAL_PAGES_ARABIC]")))
    check("header & footer", lambda: (lambda r: "ok")(run("headerFooter", [p3], pattern="[BASE_NAME] [DATE]", pageLocation="top-left")))
    check("bates numbering", lambda: (lambda r: r["task"].get("remark", ""))(run("batesNumbering", [p3, p2], style="number", batesStartFrom=1)))
    check("watermark text", lambda: (lambda r: (expect("DRAFT" in pdf_text(r["data"]), "text"), "found 'DRAFT'")[1])(run("watermark", [p3], watermarkText="DRAFT")))
    check("watermark image", lambda: (lambda r: "ok")(run("watermark", [p3, pic], watermarkType="image", pageLocation="tile")))
    # security & document
    def enc_dec():
        r = run("encrypt", [p3], userPassword="pw123")
        fid = upload("locked.pdf", r["data"])
        info = c.get(f"/api/files/{fid}").json()
        expect(info["passwordProtected"], "not protected")
        expect(c.post(f"/api/files/{fid}/password", json={"password": "wrong"}).status_code == 400, "wrong pw accepted")
        c.post(f"/api/files/{fid}/password", json={"password": "pw123"}).raise_for_status()
        r2 = run("decrypt", [fid])
        import pypdf

        expect(not pypdf.PdfReader(io.BytesIO(r2["data"])).is_encrypted, "still encrypted")
        return "locked, then unlocked"
    check("encrypt + unlock", enc_dec)
    check("edit metadata", lambda: (lambda r: "ok")(run("editMetadata", [p3], **{"metadata.Title": "Smoke", "metadata.Author": "Tester"})))
    check("edit bookmarks", lambda: (lambda r: "ok")(run("editBookmarks", [p3], bookmarks="Start | 1\n  Detail | 2\nEnd | 3")))
    check("delete annotations", lambda: (lambda r: r["task"].get("remark", ""))(run("deleteAnnotations", [linked], annotationTypes="all")))
    # convert
    check("pdf to jpg", lambda: (lambda r: f"{len(zip_names(r['data']))} images")(run("pdfToImage", [p3], imageFormat="jpg", resolution=75)))
    check("pdf to png (1 page)", lambda: (lambda r: (expect(r["data"][:4] == b"\x89PNG", "png"), r["name"])[1])(
        run("pdfToImage", [p3], imageFormat="png", pages="1", resolution=75)))
    check("jpg to pdf", lambda: (lambda r: (expect(is_pdf(r["data"]), "pdf"), r["name"])[1])(run("jpgToPdf", [pic], pageSize="a4")))
    check("extract images", lambda: (lambda r: r["task"].get("remark", ""))(run("extractImages", [img_pdf])))
    check("pdf to text", lambda: (lambda r: (expect(b"Page 2" in r["data"], "text"), r["name"])[1])(run("extractText", [p3])))
    check("pdf to excel", lambda: (lambda r: (expect(r["data"][:2] == b"PK", "xlsx"), r["name"])[1])(run("pdfToExcelNext", [tab])))
    check("pdf to csv", lambda: (lambda r: (expect(b"Apples" in r["data"], "csv"), r["name"])[1])(run("pdfToExcelNext", [tab], exportFormat="csv")))
    check("pdf to word", lambda: (lambda r: (expect(r["data"][:2] == b"PK", "docx"), r["name"])[1])(run("pdfToOfficeWord", [tab])))
    check("pdf to powerpoint", lambda: (lambda r: (expect(r["data"][:2] == b"PK", "pptx"), r["name"])[1])(run("pdfToPpt", [p3])))
    def word():
        r = run("docToPdf", [upload("letter.docx", docx())])
        expect("came from Word" in pdf_text(r["data"]), "text missing")
        engine = "basic fallback" if r["task"]["warnings"] else "LibreOffice"
        return f"via {engine}"
    check("word to pdf", word)
    def html():
        r = run("htmlToPdf", [], htmlCode="<h1>Hello from HTML</h1><p>smoke</p>")
        expect("Hello from HTML" in pdf_text(r["data"]), "text missing")
        r2 = run("htmlToPdf", [], urls="https://example.com")
        return f"html code ok, url ok ({len(r2['data'])} bytes)"
    check("html to pdf", html)
    def ocr():
        fid = upload("scan.pdf", scan_pdf())
        expect("4711" not in pdf_text(c.get(f"/api/files/{fid}/view").content), "scan already has text")
        r = run("ocr", [fid], lang="eng")
        txt = pdf_text(r["data"])
        expect("4711" in txt, f"OCR text missing: {txt[:80]!r}")
        return "recognized 'Scanned Invoice 4711'"
    check("ocr", ocr)
    # editor
    def editor():
        r = run("edit", [p3], pageOrder=["1:0", "-1:0", "2:90"],
                appendTextOperations=[{"text": "Edited by smoke test", "position": {"x": 72, "y": 400}, "fontSize": 16, "pages": "1"}],
                shapeOperations=[{"pages": "1", "position": {"x": 300, "y": 200}, "width": 80, "height": 40, "shape": "rectangle", "borderWidth": 2, "borderColor": "#ff0000"}],
                appendLinkOperations=[{"page": "1", "href": "example.com", "boundingBoxes": [{"left": 72, "right": 200, "top": 120, "bottom": 100}]}])
        expect("Edited by smoke test" in pdf_text(r["data"]) and pdf_pages(r["data"]) == 3, "edit not applied")
        return "text, shape, link, blank page, rotation"
    check("pdf editor", editor)
    def fill():
        fid = upload("form.pdf", form_pdf())
        r = run("fillSign", [fid], formValues={"name": "Ada Lovelace", "agree": True},
                imageOperations=[{"pages": "1", "position": {"x": 350, "y": 100}, "width": 120, "height": 40,
                                  "image": "data:image/png;base64," + __import__("base64").b64encode(png(240, 80)).decode()}])
        import pypdf

        f = pypdf.PdfReader(io.BytesIO(r["data"])).get_fields()
        expect(f["name"]["/V"] == "Ada Lovelace" and f["agree"]["/V"] == "/Yes", "values not filled")
        return "form filled + signature image"
    check("fill & sign", fill)
    # platform
    def platform():
        r = c.post("/api/files", json={"url": "http://127.0.0.1/x.pdf"})
        expect(r.status_code == 400, "private URL not blocked")
        r = c.post("/api/files", json={"url": "https://www.w3.org/WAI/ER/tests/xhtml/testfiles/resources/pdf/dummy.pdf"})
        expect(r.status_code == 200 and r.json()["pages"] == 1, f"url import failed: {r.text[:100]}")
        data = pdf(2)
        h = {"X-Upload-Id": "smoketest1234", "X-File-Name": "chunked.pdf"}
        half = len(data) // 2
        c.post("/api/files/upload-chunked", files={"file": ("c.pdf", data[:half])}, headers={**h, "Content-Range": f"bytes 0-{half - 1}/{len(data)}"})
        r = c.post("/api/files/upload-chunked", files={"file": ("c.pdf", data[half:])}, headers={**h, "Content-Range": f"bytes {half}-{len(data) - 1}/{len(data)}"})
        expect(r.json().get("pages") == 2, "chunked upload failed")
        res = run("extractPages", [p3], pages="1")
        cont = c.post(f"/api/tasks/{res['task']['id']}/continue").json()
        expect(cont[0]["pages"] == 1, "continue failed")
        bad = c.get(f"/api/tasks/{res['task']['id']}/download/x.pdf?s=web-wrong")
        expect(bad.status_code == 403, "bad signature accepted")
        return "URL import, SSRF block, chunked upload, continue, signed links"
    check("platform features", platform)

    failed = [r for r in results if not r[0]]
    print(f"\n{len(results) - len(failed)}/{len(results)} passed in {sum(r[3] for r in results):.1f}s")
    for _, label, note, _ in failed:
        print(f"  FAILED: {label}: {note}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
