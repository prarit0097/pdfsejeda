import io
import zipfile

import pytest
from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.tools.convert import tesseract_path
from conftest import make_pdf, make_png, open_out, run_tool, store


def _table_pdf():
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, 780, "Quarterly report")
    c.setFont("Helvetica", 11)
    c.drawString(72, 750, "Revenue grew in every region during the quarter.")
    rows = [["Region", "Q1", "Q2"], ["North", "1200", "1350"], ["South", "980", "1010"]]
    x0, y0, cw, rh = 72, 700, 120, 22
    for r, row in enumerate(rows):
        for col, v in enumerate(row):
            c.rect(x0 + col * cw, y0 - r * rh, cw, rh)
            c.drawString(x0 + col * cw + 6, y0 - r * rh + 7, v)
    c.showPage()
    c.save()
    return buf.getvalue()


def test_pdf_to_images(pdf3):
    task, out = run_tool("pdfToImage", [pdf3], imageFormat="png", resolution=75, pages="1-2")
    assert sorted(n for n in out if n.endswith(".png")) == ["three-1.png", "three-2.png"]
    task, out = run_tool("pdfToImage", [pdf3], imageFormat="tiff", tiffSingleFileWithAllPages=True, imageColorType="black_white")
    assert task.result_filename == "three.tif"


def test_images_to_pdf():
    a = store(make_png(400, 200), "a.png", "image/png")
    b = store(make_png(100, 300), "b.png", "image/png")
    task, out = run_tool("jpgToPdf", [a, b], pageSize="a4", pageMargin="small", rotations=[90, 0])
    pdf = open_out(out[task.result_filename])
    assert len(pdf.pages) == 2


def test_extract_images_and_text():
    from test_modify import _image_pdf

    f = store(_image_pdf(), "pics.pdf")
    task, out = run_tool("extractImages", [f])
    assert task.remark == "Extracted 1 image"
    t = store(make_pdf(2), "t.pdf")
    _, out = run_tool("extractText", [t])
    assert "Page 2" in out["t.txt"].decode()


def test_pdf_to_excel_and_csv():
    f = store(_table_pdf(), "report.pdf")
    _, out = run_tool("pdfToExcelNext", [f])
    ws = load_workbook(io.BytesIO(out["report.xlsx"])).active
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert ["North", 1200, 1350] in values
    _, out = run_tool("pdfToExcelNext", [f], exportFormat="csv")
    assert "South,980,1010" in out["report.csv"].decode("utf-8-sig")


def test_pdf_to_word():
    from docx import Document

    f = store(_table_pdf(), "report.pdf")
    _, out = run_tool("pdfToOfficeWord", [f])
    d = Document(io.BytesIO(out["report.docx"]))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Quarterly report" in text and "Revenue grew" in text
    assert d.tables and d.tables[0].cell(1, 0).text == "North"


def test_pdf_to_ppt(pdf3):
    from pptx import Presentation

    _, out = run_tool("pdfToPpt", [pdf3])
    prs = Presentation(io.BytesIO(out["three.pptx"]))
    assert len(prs.slides) == 3


def test_word_to_pdf_fallback():
    from docx import Document

    d = Document()
    d.add_heading("Hello", 1)
    d.add_paragraph("Body text")
    buf = io.BytesIO()
    d.save(buf)
    f = store(buf.getvalue(), "letter.docx", "application/octet-stream")
    task, out = run_tool("docToPdf", [f])
    import pypdf

    assert "Body text" in pypdf.PdfReader(io.BytesIO(out["letter.pdf"])).pages[0].extract_text()


@pytest.mark.skipif(not tesseract_path(), reason="Tesseract not installed")
def test_ocr(pdf3):
    from app.storage import files

    _, out = run_tool("pdfToImage", [pdf3], imageFormat="png", pages="1", resolution=150)
    img = store(out["three-1.png"], "scan.png", "image/png")
    _, out = run_tool("jpgToPdf", [img])
    scan = store(out["scan.pdf"], "scan.pdf")
    _, out = run_tool("ocr", [scan], lang="eng")
    import pypdf

    assert "Page" in pypdf.PdfReader(io.BytesIO(out["scan.pdf"])).pages[0].extract_text()
