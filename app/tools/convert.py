"""Convert tools: PDF -> images/text/Excel/Word/PowerPoint, images -> PDF,
extract images, Word -> PDF (LibreOffice), HTML -> PDF (Chromium), OCR (Tesseract)."""
from __future__ import annotations

import csv
import io
import os
import re
import shutil
import subprocess
from pathlib import Path

import pikepdf

from ..jobs import TaskError
from ..security import UnsafeUrl, _check_host
from .common import PAGE_SIZES_PT, open_pdf, output_name, pages_or_error, pdfium_doc
from .imaging import image_to_pdf
from .registry import IMAGES, Opt, Tool, register


def _num(v, default=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def find_binary(names: list[str], extra: list[str] = ()) -> str | None:
    for n in names:
        p = shutil.which(n)
        if p:
            return p
    for p in extra:
        if os.path.exists(p):
            return p
    return None


def soffice_path():
    return find_binary(["soffice", "libreoffice"], [r"C:\Program Files\LibreOffice\program\soffice.exe",
                                                    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                                                    "/Applications/LibreOffice.app/Contents/MacOS/soffice"])


def tesseract_path():
    return find_binary(["tesseract"], [r"C:\Program Files\Tesseract-OCR\tesseract.exe", "/opt/homebrew/bin/tesseract"])


# ---------------------------------------------------------------------------
# PDF -> images


def pdf_to_image(ctx):
    from PIL import Image

    fmt = (ctx.param("imageFormat") or "jpg").lower()
    dpi = int(_num(ctx.param("resolution"), 150) or 150)
    dpi = max(36, min(dpi, 600))
    color = ctx.param("imageColorType", "rgb")
    single_tiff = fmt == "tiff" and ctx.flag("tiffSingleFileWithAllPages")
    pages_spec = "1-" if ctx.flag("allPages") else ctx.param("pages")
    for inp in ctx.inputs:
        frames = []
        with pdfium_doc(inp) as doc:
            sel = pages_or_error(pages_spec, len(doc))
            for p in sel:
                ctx.check_cancelled()
                pil = doc[p - 1].render(scale=dpi / 72, draw_annots=True, may_draw_forms=True).to_pil()
                pil = pil.convert("L") if color == "grayscale" else pil.convert("1") if color == "black_white" else pil.convert("RGB")
                if single_tiff:
                    frames.append(pil)
                    continue
                ext = {"jpg": "jpg", "png": "png", "tiff": "tif"}.get(fmt, "jpg")
                out = ctx.output(f"{inp.stem}-{p}.{ext}")
                if ext == "jpg":
                    (pil if pil.mode != "1" else pil.convert("L")).save(out, "JPEG", quality=90, dpi=(dpi, dpi))
                elif ext == "png":
                    pil.save(out, "PNG", optimize=True, dpi=(dpi, dpi))
                else:
                    pil.save(out, "TIFF", compression="group4" if pil.mode == "1" else "tiff_lzw", dpi=(dpi, dpi))
        if single_tiff and frames:
            frames[0].save(ctx.output(f"{inp.stem}.tif"), "TIFF", save_all=True, append_images=frames[1:],
                           compression="group4" if color == "black_white" else "tiff_lzw", dpi=(dpi, dpi))
    ctx.remark("Pages converted to images")


register(Tool(
    type="pdfToImage", slug="pdf-to-jpg", title="PDF to JPG", icon="image", category="from_pdf", popular=True,
    tagline="Turn each page into a JPG, PNG or TIFF image.",
    handler=pdf_to_image, page_ui="select", keywords="pdf to image png tiff jpeg picture convert",
    options=[
        Opt("pages", "Pages", "pages", placeholder="all pages"),
        Opt("imageFormat", "Format", "radio", [("jpg", "JPG"), ("png", "PNG"), ("tiff", "TIFF")], "jpg"),
        Opt("resolution", "Resolution", "radio", [("75", "75 dpi"), ("150", "150 dpi"), ("220", "220 dpi"), ("300", "300 dpi")], "150"),
        Opt("imageColorType", "Colors", "radio", [("rgb", "Color"), ("grayscale", "Grayscale"), ("black_white", "Black & white")], "rgb", advanced=True),
        Opt("tiffSingleFileWithAllPages", "One multi-page TIFF instead of one file per page", "checkbox", show_if={"imageFormat": "tiff"}),
    ],
))


# ---------------------------------------------------------------------------
# images -> PDF


def jpg_to_pdf(ctx):
    size = (ctx.param("pageSize") or "auto").lower()
    orientation = ctx.param("pageOrientation", "auto")
    margin = {"none": 0, "small": 20, "big": 50}.get(str(ctx.param("pageMargin", "none")), _num(ctx.param("pageMargin"), 0))
    custom = None
    if size == "custom":
        unit = 72 if ctx.param("unit", "inch") == "inch" else 72 / 2.54
        custom = (_num(ctx.param("pageSizeWidth")) * unit, _num(ctx.param("pageSizeHeight")) * unit)
        if custom[0] < 36 or custom[1] < 36:
            raise TaskError("Enter a custom page width and height")
    rotations = ctx.params.get("rotations") or []
    merge = ctx.flag("mergeOutputs", True)
    combined = pikepdf.new()
    for n, inp in enumerate(ctx.inputs):
        ctx.check_cancelled()
        rot = int(_num(rotations[n] if n < len(rotations) else 0))
        if inp.name.lower().endswith(".pdf"):
            part = open_pdf(inp)
        else:
            part = image_to_pdf(inp.path, size, orientation, margin, custom, rotate=rot)
        if merge:
            combined.pages.extend(part.pages)
        else:
            part.save(ctx.output(f"{inp.stem}.pdf"))
    if merge:
        name = f"{ctx.inputs[0].stem}.pdf" if len(ctx.inputs) == 1 else "images.pdf"
        combined.save(ctx.output(name))
    ctx.remark(f"Converted {len(ctx.inputs)} image{'s' if len(ctx.inputs) != 1 else ''}")


register(Tool(
    type="jpgToPdf", slug="jpg-to-pdf", title="JPG to PDF", icon="images", category="to_pdf", popular=True,
    tagline="Turn photos and images (JPG, PNG, GIF, TIFF, BMP, WEBP) into a PDF.",
    handler=jpg_to_pdf, accept=IMAGES + ["pdf"], page_ui="files", keywords="image to pdf png photo picture convert",
    options=[
        Opt("pageSize", "Page size", "select", [("auto", "Same as image"), ("a4", "A4"), ("letter", "Letter"), ("legal", "Legal"),
            ("a3", "A3"), ("a5", "A5"), ("custom", "Custom")], "auto"),
        Opt("pageOrientation", "Orientation", "radio", [("auto", "Automatic"), ("portrait", "Portrait"), ("landscape", "Landscape")], "auto"),
        Opt("pageMargin", "Margin", "radio", [("none", "None"), ("small", "Small"), ("big", "Big")], "none"),
        Opt("unit", "Unit", "radio", [("inch", "inch"), ("cm", "cm")], "inch", show_if={"pageSize": "custom"}),
        Opt("pageSizeWidth", "Width", "number", step=0.1, show_if={"pageSize": "custom"}),
        Opt("pageSizeHeight", "Height", "number", step=0.1, show_if={"pageSize": "custom"}),
        Opt("mergeOutputs", "Put all images in one PDF", "checkbox", default=True, advanced=True),
    ],
))


# ---------------------------------------------------------------------------
# extract images / text


def extract_images(ctx):
    from pikepdf import PdfImage

    from .modify import iter_image_xobjects

    count = 0
    for inp in ctx.inputs:
        pdf = open_pdf(inp)
        for n, (xo, _) in enumerate(iter_image_xobjects(pdf), 1):
            ctx.check_cancelled()
            try:
                img = PdfImage(xo)
                if img.width < 16 or img.height < 16:
                    continue
                target = ctx.output(f"{inp.stem}-image-{n}")
                img.extract_to(fileprefix=str(target))
                count += 1
            except Exception:
                try:
                    PdfImage(xo).as_pil_image().convert("RGB").save(ctx.output(f"{inp.stem}-image-{n}.png"))
                    count += 1
                except Exception:
                    continue
    if not count:
        raise TaskError("No images were found in this PDF")
    ctx.remark(f"Extracted {count} image{'s' if count != 1 else ''}")


register(Tool(
    type="extractImages", slug="extract-images-from-pdf", title="Extract images", icon="images", category="from_pdf",
    tagline="Save every picture embedded in a PDF as a separate image file.",
    handler=extract_images, keywords="images pictures photos save export",
))


def extract_text(ctx):
    for inp in ctx.inputs:
        parts = []
        with pdfium_doc(inp) as doc:
            sel = pages_or_error(ctx.param("pages"), len(doc))
            for p in sel:
                tp = doc[p - 1].get_textpage()
                parts.append(tp.get_text_bounded().replace("\r\n", "\n").strip())
                tp.close()
        text = "\n\n\f".join(parts) if ctx.flag("pageBreaks") else "\n\n".join(parts)
        if not text.strip():
            ctx.warn(f"'{inp.name}' contains no selectable text. If it is a scan, run OCR first.")
        ctx.output(f"{inp.stem}.txt").write_text(text, encoding="utf-8")
    ctx.remark("Text extracted")


register(Tool(
    type="extractText", slug="extract-text-from-pdf", title="PDF to text", icon="text", category="from_pdf",
    tagline="Copy all the text out of a PDF into a plain .txt file.",
    handler=extract_text, keywords="text txt copy extract",
    options=[
        Opt("pages", "Pages", "pages", placeholder="all pages", advanced=True),
        Opt("pageBreaks", "Mark page breaks with a form-feed character", "checkbox", advanced=True),
    ],
))


# ---------------------------------------------------------------------------
# PDF -> Excel / CSV


def _tables_for(inp, sel):
    import pdfplumber

    found = []
    with pdfplumber.open(str(inp.path), password=inp.password or "") as doc:
        for p in sel:
            page = doc.pages[p - 1]
            tables = page.extract_tables() or page.extract_tables({"vertical_strategy": "text", "horizontal_strategy": "text"})
            for t in tables:
                rows = [[(c or "").strip() for c in row] for row in t if any((c or "").strip() for c in row)]
                if rows and max(len(r) for r in rows) > 1:
                    found.append((p, rows))
    return found


def _cell(v: str):
    s = v.replace("\n", " ").strip()
    if re.fullmatch(r"-?\d{1,3}(,\d{3})*(\.\d+)?|-?\d+(\.\d+)?", s):
        try:
            return float(s.replace(",", "")) if "." in s else int(s.replace(",", ""))
        except ValueError:
            return s
    return s


def pdf_to_excel(ctx):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    fmt = ctx.param("exportFormat", "excel")
    single = ctx.flag("singleSheet")
    for inp in ctx.inputs:
        with pdfium_doc(inp) as doc:
            total = len(doc)
        tables = _tables_for(inp, pages_or_error(ctx.param("pages"), total))
        if not tables:
            raise TaskError(f"No tables were found in '{inp.name}'")
        if fmt == "csv":
            groups = [("all", [r for _, t in tables for r in t + [[]]])] if single else [(f"p{p}-t{i + 1}", t) for i, (p, t) in enumerate(tables)]
            for name, rows in groups:
                buf = io.StringIO()
                csv.writer(buf).writerows(rows)
                suffix = "" if single or len(groups) == 1 else f"-{name}"
                ctx.output(f"{inp.stem}{suffix}.csv").write_text(buf.getvalue(), encoding="utf-8-sig")
            continue
        wb = Workbook()
        wb.remove(wb.active)
        ws = None
        for i, (p, rows) in enumerate(tables):
            if ws is None or not single:
                ws = wb.create_sheet(("Tables" if single else f"Page {p} table {i + 1}")[:31])
                r0 = 1
            else:
                r0 = ws.max_row + 2
            for r, row in enumerate(rows):
                for c, v in enumerate(row):
                    cell = ws.cell(row=r0 + r, column=c + 1, value=_cell(v))
                    if r == 0:
                        cell.font = Font(bold=True)
            for c in range(1, max(len(x) for x in rows) + 1):
                width = max((len(str(x[c - 1])) for x in rows if len(x) >= c), default=8)
                ws.column_dimensions[get_column_letter(c)].width = min(60, max(8, width + 2))
        wb.save(ctx.output(f"{inp.stem}.xlsx"))
    ctx.remark(f"Found {len(tables)} table{'s' if len(tables) != 1 else ''}")


register(Tool(
    type="pdfToExcelNext", slug="pdf-to-excel", title="PDF to Excel", icon="table", category="from_pdf", popular=True,
    tagline="Pull tables out of a PDF into an Excel workbook or CSV files.",
    handler=pdf_to_excel, keywords="xlsx csv table spreadsheet convert",
    options=[
        Opt("exportFormat", "Format", "radio", [("excel", "Excel (.xlsx)"), ("csv", "CSV")], "excel"),
        Opt("singleSheet", "Put all tables in one sheet / file", "checkbox", advanced=True),
        Opt("pages", "Pages", "pages", placeholder="all pages", advanced=True),
    ],
))


# ---------------------------------------------------------------------------
# PDF -> Word


def pdf_to_word(ctx):
    import pdfplumber
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    for inp in ctx.inputs:
        document = Document()
        style = document.styles["Normal"]
        style.font.size = Pt(11)
        with pdfplumber.open(str(inp.path), password=inp.password or "") as doc, pdfium_doc(inp) as pdoc:
            body_sizes = []
            for page in doc.pages[:5]:
                body_sizes += [round(c["size"]) for c in page.chars[:2000]]
            body = max(set(body_sizes), key=body_sizes.count) if body_sizes else 11
            for pi, page in enumerate(doc.pages):
                ctx.check_cancelled()
                if pi:
                    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                blocks = []  # (top, kind, payload)
                tables = page.find_tables()
                boxes = [t.bbox for t in tables]
                for t in tables:
                    blocks.append((t.bbox[1], "table", t.extract()))
                for img in page.images:
                    w, h = img["x1"] - img["x0"], img["bottom"] - img["top"]
                    if w > 40 and h > 40:
                        blocks.append((img["top"], "image", (img["x0"], img["top"], img["x1"], img["bottom"])))
                words = [w for w in page.extract_words(extra_attrs=["size", "fontname"], keep_blank_chars=False, use_text_flow=True)
                         if not any(b[0] <= w["x0"] and w["x1"] <= b[2] and b[1] <= w["top"] and w["bottom"] <= b[3] for b in boxes)]
                for para in _paragraphs(words):
                    blocks.append((para["top"], "text", para))
                blocks.sort(key=lambda b: b[0])
                for _, kind, payload in blocks:
                    if kind == "table":
                        rows = [r for r in payload if r]
                        if not rows:
                            continue
                        ncols = max(len(r) for r in rows)
                        table = document.add_table(rows=len(rows), cols=ncols)
                        table.style = "Table Grid"
                        for r, row in enumerate(rows):
                            for c, v in enumerate(row):
                                table.cell(r, c).text = (v or "").strip()
                    elif kind == "image":
                        x0, top, x1, bottom = payload
                        scale = 2
                        pil = pdoc[pi].render(scale=scale, crop=(x0, float(page.height) - bottom, float(page.width) - x1, top)).to_pil()
                        buf = io.BytesIO()
                        pil.convert("RGB").save(buf, "PNG")
                        buf.seek(0)
                        document.add_picture(buf, width=Pt(min(x1 - x0, 450)))
                    else:
                        size = payload["size"]
                        text = payload["text"]
                        if size >= body + 5 and len(text) < 120:
                            document.add_heading(text, level=1)
                        elif size >= body + 2 and len(text) < 160:
                            document.add_heading(text, level=2)
                        else:
                            p = document.add_paragraph()
                            run = p.add_run(text)
                            run.bold = payload["bold"]
                            if abs(size - body) >= 1:
                                run.font.size = Pt(size)
        document.save(ctx.output(f"{inp.stem}.docx"))
    ctx.remark("Converted to Word")


def _paragraphs(words):
    """Group words into lines, then lines into paragraphs by vertical gaps."""
    lines = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if lines and abs(lines[-1]["top"] - w["top"]) < max(2.5, w["size"] * 0.35):
            lines[-1]["words"].append(w)
        else:
            lines.append({"top": w["top"], "bottom": w["bottom"], "words": [w]})
        lines[-1]["bottom"] = max(lines[-1]["bottom"], w["bottom"])
    paras = []
    for ln in lines:
        text = " ".join(x["text"] for x in sorted(ln["words"], key=lambda x: x["x0"]))
        size = max(x["size"] for x in ln["words"])
        bold = all("bold" in x["fontname"].lower() for x in ln["words"])
        prev = paras[-1] if paras else None
        if prev and ln["top"] - prev["bottom"] < size * 0.9 and abs(prev["size"] - size) < 1 and prev["bold"] == bold:
            if prev["text"].endswith("-"):  # word broken across lines
                prev["text"] = prev["text"][:-1] + text
            else:
                prev["text"] += " " + text
            prev["bottom"] = ln["bottom"]
        else:
            paras.append({"top": ln["top"], "bottom": ln["bottom"], "text": text, "size": size, "bold": bold})
    return paras


register(Tool(
    type="pdfToOfficeWord", slug="pdf-to-word", title="PDF to Word", icon="doc", category="from_pdf", popular=True,
    tagline="Convert a PDF into an editable Word document, keeping paragraphs, headings, tables and pictures.",
    handler=pdf_to_word, keywords="docx doc word editable convert",
))


# ---------------------------------------------------------------------------
# PDF -> PowerPoint


def pdf_to_ppt(ctx):
    from pptx import Presentation
    from pptx.util import Emu

    for inp in ctx.inputs:
        prs = Presentation()
        with pdfium_doc(inp) as doc:
            w, h = doc[0].get_size()
            prs.slide_width, prs.slide_height = Emu(int(w * 12700)), Emu(int(h * 12700))
            blank = prs.slide_layouts[6]
            for i in range(len(doc)):
                ctx.check_cancelled()
                page = doc[i]
                pil = page.render(scale=2).to_pil().convert("RGB")
                buf = io.BytesIO()
                pil.save(buf, "JPEG", quality=88)
                buf.seek(0)
                slide = prs.slides.add_slide(blank)
                slide.shapes.add_picture(buf, 0, 0, prs.slide_width, prs.slide_height)
                tp = page.get_textpage()
                text = tp.get_text_bounded().strip()
                tp.close()
                if text:
                    slide.notes_slide.notes_text_frame.text = text
        prs.save(ctx.output(f"{inp.stem}.pptx"))
    ctx.remark("Converted to PowerPoint (one slide per page; page text is in the speaker notes)")


register(Tool(
    type="pdfToPpt", slug="pdf-to-ppt", title="PDF to PowerPoint", icon="slides", category="from_pdf",
    tagline="Turn each PDF page into a PowerPoint slide.",
    handler=pdf_to_ppt, keywords="pptx powerpoint slides presentation convert",
))


# ---------------------------------------------------------------------------
# Word -> PDF


def doc_to_pdf(ctx):
    soffice = soffice_path()
    for inp in ctx.inputs:
        ctx.check_cancelled()
        if soffice:
            src = ctx.temp(inp.name)
            shutil.copyfile(inp.path, src)
            profile = ctx.temp("lo-profile")
            cmd = [soffice, f"-env:UserInstallation=file:///{str(profile).replace(os.sep, '/')}", "--headless",
                   "--convert-to", "pdf", "--outdir", str(ctx.work_dir), str(src)]
            try:
                subprocess.run(cmd, check=True, timeout=180, capture_output=True)
            except (subprocess.SubprocessError, OSError):
                raise TaskError(f"Could not convert '{inp.name}'") from None
            produced = ctx.work_dir / (Path(src).stem + ".pdf")
            if not produced.exists():
                raise TaskError(f"Could not convert '{inp.name}'")
            shutil.move(produced, ctx.output(f"{inp.stem}.pdf"))
        elif inp.name.lower().endswith((".docx", ".txt")):
            _simple_docx_to_pdf(inp, ctx.output(f"{inp.stem}.pdf"))
            ctx.warn("Converted with the basic built-in converter (text and tables only). Install LibreOffice on the server for exact layout.")
        else:
            raise TaskError("Converting this file type needs LibreOffice on the server")
    ctx.remark("Converted to PDF")


def _simple_docx_to_pdf(inp, out: Path):
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from xml.sax.saxutils import escape

    from .draw import unicode_font

    styles = getSampleStyleSheet()
    uni = unicode_font()
    if uni:
        for s in styles.byName.values():
            s.fontName = uni
    story = []
    if inp.name.lower().endswith(".txt"):
        for line in inp.path.read_text(encoding="utf-8", errors="replace").splitlines():
            story.append(Paragraph(escape(line) or "&nbsp;", styles["Normal"]))
    else:
        from docx import Document

        d = Document(str(inp.path))
        body = d.element.body
        for child in body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                from docx.text.paragraph import Paragraph as DP

                p = DP(child, d)
                name = (p.style.name or "").lower() if p.style is not None else ""
                st = styles["Heading1"] if name.startswith("heading 1") or name == "title" else styles["Heading2"] if name.startswith("heading") else styles["Normal"]
                html = "".join(
                    (f"<b>{escape(r.text)}</b>" if r.bold else f"<i>{escape(r.text)}</i>" if r.italic else escape(r.text)) for r in p.runs)
                story.append(Paragraph(html or "&nbsp;", st))
            elif tag == "tbl":
                from docx.table import Table as DT

                t = DT(child, d)
                data = [[Paragraph(escape(c.text), styles["Normal"]) for c in row.cells] for row in t.rows]
                if data:
                    tbl = Table(data)
                    tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                    story += [tbl, Spacer(1, 8)]
    SimpleDocTemplate(str(out), pagesize=PAGE_SIZES_PT["a4"], leftMargin=56, rightMargin=56, topMargin=56, bottomMargin=56).build(story or [Spacer(1, 1)])


register(Tool(
    type="docToPdf", slug="word-to-pdf", title="Word to PDF", icon="doc", category="to_pdf", popular=True,
    tagline="Convert Word documents (DOC, DOCX, ODT, RTF) into PDF.",
    handler=doc_to_pdf, accept=["doc", "docx", "odt", "rtf", "txt"], keywords="docx doc word convert",
    requires="LibreOffice (for exact layout)",
))


# ---------------------------------------------------------------------------
# HTML -> PDF


def html_to_pdf(ctx):
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        raise TaskError("HTML to PDF needs Chromium on the server (pip install playwright && playwright install chromium)") from None
    urls = [u.strip() for u in re.split(r"[\n,]+", ctx.param("urls") or "") if u.strip()]
    html_code = ctx.param("htmlCode") or ""
    html_files = [i for i in ctx.inputs if i.name.lower().endswith((".html", ".htm"))]
    if not urls and not html_code and not html_files:
        raise TaskError("Enter a web address or some HTML")
    for i, u in enumerate(urls):
        if "://" not in u:
            urls[i] = u = "https://" + u
        try:
            _check_host(u)
        except UnsafeUrl as e:
            raise TaskError(f"{u}: {e}") from None
    size = ctx.param("pageSize", "A4")
    landscape = ctx.param("pageOrientation", "auto") == "landscape"
    margin = f"{_num(ctx.param('pageMargin'), 0)}{ctx.param('pageMarginUnits', 'px')}"
    delay = min(5, max(0, int(_num(ctx.param("delay"), 0)))) * 1000
    viewport = int(_num(ctx.param("viewportWidth"), 1280) or 1280)
    media = "print" if ctx.flag("usePrintMedia") else "screen"

    def guard(route):
        try:
            _check_host(route.request.url) if route.request.url.startswith("http") else None
            route.continue_()
        except UnsafeUrl:
            route.abort()

    with sync_playwright() as p:
        browser = p.chromium.launch()
        try:
            page = browser.new_page(viewport={"width": max(320, min(viewport, 3840)), "height": 900})
            page.route("**/*", guard)
            jobs = [(u, None) for u in urls]
            if html_code:
                jobs.append((None, html_code))
            jobs += [(None, f.path.read_text(encoding="utf-8", errors="replace")) for f in html_files]
            for n, (url, html) in enumerate(jobs, 1):
                ctx.check_cancelled()
                if url:
                    page.goto(url, wait_until="networkidle", timeout=60000)
                    name = re.sub(r"[^\w.-]+", "_", url.split("://", 1)[1])[:80]
                else:
                    page.set_content(html, wait_until="networkidle", timeout=60000)
                    name = f"page-{n}"
                if delay:
                    page.wait_for_timeout(delay)
                page.emulate_media(media=media)
                page.pdf(path=str(ctx.output(name + ".pdf")), format=size, landscape=landscape, print_background=True,
                         margin={"top": margin, "bottom": margin, "left": margin, "right": margin})
        finally:
            browser.close()
    ctx.remark("Converted to PDF")


register(Tool(
    type="htmlToPdf", slug="html-to-pdf", title="HTML to PDF", icon="globe", category="to_pdf",
    tagline="Save web pages or HTML code as PDF documents.",
    handler=html_to_pdf, accept=["html", "htm"], min_files=0, keywords="web page website url save print",
    requires="Chromium via Playwright",
    options=[
        Opt("urls", "Web addresses (one per line)", "textarea", placeholder="https://example.com"),
        Opt("htmlCode", "…or paste HTML", "textarea", placeholder="<h1>Hello</h1>", advanced=True),
        Opt("pageSize", "Page size", "select", [("A4", "A4"), ("Letter", "Letter"), ("Legal", "Legal"), ("A3", "A3")], "A4"),
        Opt("pageOrientation", "Orientation", "radio", [("portrait", "Portrait"), ("landscape", "Landscape")], "portrait"),
        Opt("viewportWidth", "Browser width (px)", "number", default=1280, min=320, max=3840, advanced=True),
        Opt("pageMargin", "Margin", "number", default=0, min=0, advanced=True),
        Opt("pageMarginUnits", "Margin unit", "radio", [("px", "px"), ("mm", "mm"), ("in", "in")], "px", advanced=True),
        Opt("delay", "Wait before saving (seconds)", "select", [(str(i), str(i)) for i in range(6)], "0", advanced=True),
        Opt("usePrintMedia", "Use the page's print styles", "checkbox", advanced=True),
    ],
))


# ---------------------------------------------------------------------------
# OCR

OCR_LANGS = [("eng", "English"), ("hin", "Hindi"), ("spa", "Spanish"), ("fra", "French"), ("deu", "German"), ("ita", "Italian"),
             ("por", "Portuguese"), ("nld", "Dutch"), ("rus", "Russian"), ("ara", "Arabic"), ("chi_sim", "Chinese (simplified)"),
             ("jpn", "Japanese"), ("kor", "Korean"), ("tur", "Turkish"), ("pol", "Polish"), ("ben", "Bengali"), ("tam", "Tamil"),
             ("tel", "Telugu"), ("mar", "Marathi"), ("guj", "Gujarati"), ("urd", "Urdu")]


def ocr(ctx):
    from .draw import draw_form, invert, user_to_visible

    tess = tesseract_path()
    if not tess:
        raise TaskError("OCR needs Tesseract on the server (apt install tesseract-ocr)")
    lang = re.sub(r"[^a-z_+]", "", str(ctx.param("lang") or "eng")) or "eng"
    formats = ctx.params.get("outputFormats") or ["pdf"]
    formats = [formats] if isinstance(formats, str) else formats
    dpi = 300
    for inp in ctx.inputs:
        pdf = open_pdf(inp)
        texts = []
        with pdfium_doc(inp) as doc:
            sel = pages_or_error(ctx.param("pages"), len(doc))
            for p in sel:
                ctx.check_cancelled()
                img = ctx.temp(f"p{p}.png")
                doc[p - 1].render(scale=dpi / 72).to_pil().convert("L").save(img, dpi=(dpi, dpi))
                base = ctx.temp(f"p{p}")
                try:
                    subprocess.run([tess, str(img), str(base), "-l", lang, "--dpi", str(dpi), "-c", "textonly_pdf=1", "pdf", "txt"],
                                   check=True, timeout=300, capture_output=True)
                except subprocess.CalledProcessError as e:
                    msg = e.stderr.decode(errors="replace")
                    if "Failed loading language" in msg or "Error opening data file" in msg:
                        raise TaskError(f"The OCR language '{lang}' is not installed on the server") from None
                    raise TaskError("OCR failed on page %d" % p) from None
                except (subprocess.SubprocessError, OSError):
                    raise TaskError("OCR failed on page %d" % p) from None
                texts.append(Path(str(base) + ".txt").read_text(encoding="utf-8", errors="replace"))
                with pikepdf.open(str(base) + ".pdf") as layer:
                    lp = pdf.copy_foreign(layer.pages[0].obj)
                    form = pikepdf.Page(lp).as_form_xobject()
                    lw, lh = (float(v) for v in layer.pages[0].mediabox[2:])
                page = pdf.pages[p - 1]
                from .draw import visible_size

                vw, vh = visible_size(page)
                m = (vw / lw, 0, 0, vh / lh, 0, 0)
                from .draw import mul

                draw_form(pdf, page, form, mul(m, invert(user_to_visible(page))))
        if "pdf" in formats:
            pdf.save(ctx.output(output_name(ctx, inp)))
        if "txt" in formats:
            ctx.output(f"{inp.stem}.txt").write_text("\n\f".join(texts), encoding="utf-8")
    ctx.remark("Text recognized — the PDF is now searchable")


register(Tool(
    type="ocr", slug="ocr-pdf", title="OCR PDF", icon="scan", category="optimize",
    tagline="Recognize the text in scanned PDFs so you can search, select and copy it.",
    handler=ocr, keywords="scan recognize searchable text tesseract",
    requires="Tesseract OCR",
    extra_categories=["from_pdf"],
    options=[
        Opt("lang", "Document language", "select", OCR_LANGS, "eng"),
        Opt("outputFormats", "Output", "radio", [("pdf", "Searchable PDF"), ("txt", "Text file")], "pdf"),
        Opt("pages", "Pages", "pages", placeholder="all pages", advanced=True),
    ],
))
