"""Explicit canonical tool scenarios shared by API and browser regression tests.

All documents are generated. Output validators inspect actual document contents.
Catalogue drift is a failure: registering a tool requires adding its scenario.
"""
import base64
from functools import lru_cache
import io
import zipfile

import numpy as np
from PIL import Image
import pypdf
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader


@lru_cache(maxsize=16)
def fixture(kind="pdf"):
    out = io.BytesIO()
    if kind == "image":
        Image.new("RGB", (240, 120), (25, 50, 180)).save(out, "PNG")
        return "pool.png", out.getvalue(), "image/png"
    if kind == "docx":
        from docx import Document
        d = Document(); d.add_heading("POOL WORD LETTER", 1)
        d.add_paragraph("Generated document for complete regression testing.")
        d.save(out)
        return "pool.docx", out.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if kind == "scan":
        import pypdfium2
        doc = pypdfium2.PdfDocument(fixture("one")[1])
        page = doc[0]; bitmap = page.render(scale=3)
        bitmap.to_pil().convert("RGB").save(out, "PDF", resolution=216)
        bitmap.close(); page.close(); doc.close()
        return "scan.pdf", out.getvalue(), "application/pdf"
    if kind == "locked":
        reader = pypdf.PdfReader(io.BytesIO(fixture("pdf")[1]))
        writer = pypdf.PdfWriter(); writer.append(reader); writer.encrypt("pool-secret")
        writer.write(out)
        return "locked.pdf", out.getvalue(), "application/pdf"
    c = canvas.Canvas(out, pagesize=(595, 842)); c.setTitle("POOL SOURCE TITLE")
    count = 1 if kind in ("one", "table", "form") else 3
    for i in range(count):
        c.setFont("Helvetica-Bold", 24); c.drawString(72, 760, f"POOL PAGE {i+1}")
        c.setFont("Helvetica", 12); c.drawString(72, 730, "Generated regression text 4711")
        c.bookmarkPage(f"p{i}"); c.addOutlineEntry(f"Section {i+1}", f"p{i}", 0)
        if kind == "table":
            for row, values in enumerate((("Item", "Qty"), ("Apples", "3"), ("Pears", "7"))):
                for col, value in enumerate(values):
                    c.rect(72+120*col, 600-row*25, 120, 25)
                    c.drawString(78+120*col, 608-row*25, value)
        if kind == "imagepdf":
            noise = np.random.default_rng(42).integers(0, 256, (600, 600, 3), dtype=np.uint8)
            c.drawImage(ImageReader(Image.fromarray(noise)), 100, 250, 250, 250)
        if kind == "form":
            c.acroForm.textfield(name="pool_name", x=72, y=620, width=220, height=24)
        if kind == "links":
            c.linkURL("https://example.com", (72, 620, 300, 650))
        c.showPage()
    c.save()
    return kind + ".pdf", out.getvalue(), "application/pdf"


def case(task_type, files=("pdf",), params=None, **expected):
    return {"type": task_type, "files": files, "params": params or {}, "expected": expected}


SCENARIOS = [
    case("merge", ("pdf", "one"), pages=4),
    case("splitByPages", params={"splitMode":"every", "step":1}, pages=3, outputs=3),
    case("splitByBookmarks", pages=3, outputs=3),
    case("splitBySize", params={"size":1000,"sizeUnit":"KB"}, pages=3),
    case("splitByText", params={"area":{"x":60,"y":40,"width":400,"height":80}}, pages=3),
    case("splitDownTheMiddle", params={"mode":"vertical"}, pages=6),
    case("alternateMix", ("pdf","one"), pages=4),
    case("extractPages", params={"pages":"1,3"}, pages=2, text="POOL PAGE 3"),
    case("deletePages", params={"pages":"2"}, pages=2, absent="POOL PAGE 2"),
    case("combineReorder", params={"filePages":["0:3:0","0:1:90"]}, pages=2, first="POOL PAGE 3"),
    case("rotatePages", params={"rotation":90}, pages=3, rotation=90),
    case("compress", ("imagepdf",), params={"imageQuality":.65,"imageDpi":72}, pages=3),
    case("crop", params={"cropMode":"margins","unit":"pt","marginLeft":10,"marginRight":10,"marginTop":10,"marginBottom":10}, pages=3, crop_width=575),
    case("resizePages", params={"pageSize":"letter"}, pages=3, width=612),
    case("nup", params={"n":4}, pages=1),
    case("flip", params={"flipMode":"horizontal"}, pages=3),
    case("grayscale", ("imagepdf",), pages=3, gray=True),
    case("deskew", pages=3),
    case("flatten", params={"flattenMode":"all","resolution":75}, pages=3, raster=True),
    case("repair", pages=3, text="POOL PAGE 1"),
    case("headerFooter", params={"pattern":"POOL HEADER"}, pages=3, text="POOL HEADER"),
    case("pageNumbers", params={"pattern":"POOL [PAGE_ARABIC] / [TOTAL_PAGES_ARABIC]"}, pages=3, text="POOL 2 / 3"),
    case("batesNumbering", params={"style":"number","batesStartFrom":100}, pages=3, text="100"),
    case("watermark", params={"watermarkText":"POOL DRAFT"}, pages=3, text="POOL DRAFT"),
    case("encrypt", params={"userPassword":"pool-secret"}, pages=3, encrypted=True),
    case("decrypt", ("locked",), pages=3, decrypted=True),
    case("editMetadata", params={"metadata.Title":"POOL UPDATED"}, pages=3, title="POOL UPDATED"),
    case("editBookmarks", params={"bookmarks":"POOL START | 1\nPOOL END | 3"}, pages=3, outline="POOL START"),
    case("deleteAnnotations", ("links",), params={"annotationTypes":"all"}, pages=3, no_annotations=True),
    case("renameByText", params={"area":{"x":60,"y":40,"width":400,"height":80}}, pages=3),
    case("pdfToImage", params={"imageFormat":"png","resolution":75}, images=3),
    case("jpgToPdf", ("image",), params={"pageSize":"a4"}, pages=1),
    case("extractImages", ("imagepdf",), images=1),
    case("extractText", text_file="POOL PAGE 3"),
    case("pdfToExcelNext", ("table",), xlsx="Apples"),
    case("pdfToOfficeWord", ("table",), docx="POOL PAGE 1"),
    case("pdfToPpt", slides=3),
    case("docToPdf", ("docx",), text="POOL WORD LETTER"),
    case("htmlToPdf", (), {"htmlCode":"<h1>POOL HTML RESULT</h1><p>Generated fixture.</p>"}, pages=1, text="POOL HTML RESULT"),
    case("ocr", ("scan",), {"lang":"eng"}, pages=1, text="4711"),
    case("edit", ("one",), {"appendTextOperations":[{"text":"POOL EDITED", "position":{"x":72,"y":400},"pages":"1","fontSize":16}]}, pages=1, text="POOL EDITED"),
    case("fillSign", ("form",), {"formValues":{"pool_name":"Ada Pool"}}, pages=1, field_value="Ada Pool"),
    case("editForms", ("one",), {"appendFormFieldOperations":[{"page":"1","kind":"text","name":"pool_new","boundingBox":{"left":72,"bottom":400,"right":300,"top":430}}]}, pages=1, field_name="pool_new"),
]


def scenario_files(scenario):
    return [fixture(kind) for kind in scenario["files"]]


def unpack(data, name):
    if name.lower().endswith(".zip"):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            assert z.testzip() is None
            return [(n,z.read(n)) for n in z.namelist() if not n.endswith("/")]
    return [(name,data)]


def validate_output(scenario, data, name):
    expected = scenario["expected"]; outputs = unpack(data,name)
    assert outputs and all(blob for _,blob in outputs), "Empty result"
    if "outputs" in expected: assert len(outputs) == expected["outputs"]
    if "images" in expected:
        images = [Image.open(io.BytesIO(blob)) for _,blob in outputs]
        assert len(images) == expected["images"]
        assert all(im.width > 100 and im.height > 100 for im in images)
        for im in images: im.verify()
        return
    if "text_file" in expected:
        assert expected["text_file"] in data.decode("utf-8-sig"); return
    if "xlsx" in expected:
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(data))
        assert any(expected["xlsx"] == cell.value for ws in book for row in ws for cell in row); book.close(); return
    if "docx" in expected:
        from docx import Document
        d=Document(io.BytesIO(data)); assert expected["docx"] in "\n".join(p.text for p in d.paragraphs)
        assert any("Apples" in cell.text for table in d.tables for row in table.rows for cell in row.cells); return
    if "slides" in expected:
        from pptx import Presentation
        prs=Presentation(io.BytesIO(data)); assert len(prs.slides)==expected["slides"]
        assert all(len(slide.shapes)>0 for slide in prs.slides); return
    docs=[pypdf.PdfReader(io.BytesIO(blob)) for _,blob in outputs]
    for doc in docs:
        if expected.get("encrypted"): assert doc.is_encrypted; assert doc.decrypt("pool-secret")
        if expected.get("decrypted"): assert not doc.is_encrypted
    pages=[page for doc in docs for page in doc.pages]
    assert pages
    if "pages" in expected: assert len(pages)==expected["pages"], (scenario["type"],len(pages),expected)
    text="\n".join(page.extract_text() or "" for page in pages)
    if "text" in expected: assert expected["text"] in text, (scenario["type"],text[:500])
    if "absent" in expected: assert expected["absent"] not in text
    if "first" in expected: assert expected["first"] in pages[0].extract_text()
    if "rotation" in expected: assert all(int(p.get("/Rotate",0))%360==expected["rotation"] for p in pages)
    if "width" in expected: assert abs(float(pages[0].mediabox.width)-expected["width"])<1
    if "crop_width" in expected: assert abs(float(pages[0].cropbox.width)-expected["crop_width"])<1
    if "title" in expected: assert docs[0].metadata.title==expected["title"]
    if "outline" in expected: assert expected["outline"] in str(docs[0].outline)
    if expected.get("no_annotations"): assert all(not p.get("/Annots") for p in pages)
    if "field_name" in expected: assert expected["field_name"] in docs[0].get_fields()
    if "field_value" in expected: assert docs[0].get_fields()["pool_name"]["/V"]==expected["field_value"]
    if expected.get("raster"): assert pages[0].images and not text.strip()
    if expected.get("gray"):
        im=pages[0].images[0].image.convert("RGB"); a=np.asarray(im).astype(int)
        assert np.max(np.ptp(a,axis=2))<=2
